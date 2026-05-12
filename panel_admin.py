from flask import Flask, request, jsonify, session, send_from_directory
import json
import os
from functools import wraps

app = Flask(__name__, static_folder="panel_static")
app.secret_key = "floreria-candelaria-secret-2024"

CONFIG_FILE = "config.json"
ADMIN_USER = "admin"
ADMIN_PASS = "password"


def cargar_config():
    if not os.path.exists(CONFIG_FILE):
        return {}
    with open(CONFIG_FILE, encoding="utf-8") as f:
        return json.load(f)


def guardar_config(data):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return jsonify({"error": "No autorizado"}), 401
        return f(*args, **kwargs)
    return decorated


@app.route("/")
def index():
    return send_from_directory("panel_static", "index.html")


@app.route("/catalogo")
def catalogo():
    return send_from_directory("panel_static", "catalogo.html")


@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json()
    if data.get("usuario") == ADMIN_USER and data.get("password") == ADMIN_PASS:
        session["logged_in"] = True
        return jsonify({"ok": True})
    return jsonify({"ok": False, "error": "Usuario o contraseña incorrectos"}), 401


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/productos", methods=["GET"])
def get_productos():
    config = cargar_config()
    return jsonify(config.get("productos", {}))


@app.route("/api/productos", methods=["POST"])
@login_required
def update_productos():
    config = cargar_config()
    config["productos"] = request.get_json()
    guardar_config(config)
    return jsonify({"ok": True})


@app.route("/api/mensajes", methods=["GET"])
@login_required
def get_mensajes():
    config = cargar_config()
    return jsonify(config.get("mensajes", {}))


@app.route("/api/mensajes", methods=["POST"])
@login_required
def update_mensajes():
    config = cargar_config()
    config["mensajes"] = request.get_json()
    guardar_config(config)
    return jsonify({"ok": True})


@app.route("/api/negocio", methods=["GET"])
def get_negocio():
    config = cargar_config()
    return jsonify(config.get("negocio", {}))


@app.route("/api/negocio", methods=["POST"])
@login_required
def update_negocio():
    config = cargar_config()
    config["negocio"] = request.get_json()
    guardar_config(config)
    return jsonify({"ok": True})


@app.route("/api/pedidos", methods=["GET"])
@login_required
def get_pedidos():
    try:
        from openpyxl import load_workbook
        config = cargar_config()
        archivo = config.get("negocio", {}).get("archivo_pedidos", "pedidos.xlsx")
        if not os.path.exists(archivo):
            return jsonify([])
        wb = load_workbook(archivo)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        pedidos = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            pedido = dict(zip(headers, row))
            pedido["_fila"] = i
            pedidos.append(pedido)
        return jsonify(pedidos)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/pedidos/estado", methods=["POST"])
@login_required
def update_estado_pedido():
    try:
        from openpyxl import load_workbook
        data = request.get_json()
        fila = data.get("fila")
        nuevo_estado = data.get("estado")

        config = cargar_config()
        archivo = config.get("negocio", {}).get("archivo_pedidos", "pedidos.xlsx")

        if not os.path.exists(archivo):
            return jsonify({"error": "Archivo no encontrado"}), 404

        wb = load_workbook(archivo)
        ws = wb.active

        # Buscar columna Estado
        headers = [cell.value for cell in ws[1]]
        if "Estado" not in headers:
            # Agregar columna Estado si no existe
            col_estado = len(headers) + 1
            ws.cell(row=1, column=col_estado, value="Estado")
        else:
            col_estado = headers.index("Estado") + 1

        ws.cell(row=fila, column=col_estado, value=nuevo_estado)
        wb.save(archivo)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    os.makedirs("panel_static", exist_ok=True)
    app.run(port=5001, debug=True)
